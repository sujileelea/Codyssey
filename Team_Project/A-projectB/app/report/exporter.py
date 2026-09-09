"""데이터 내보내기: CSV / JSONL / Excel(xlsx). 필터(--status summarized 등) 지원."""

from __future__ import annotations

import csv
import json
import logging
from datetime import datetime
from pathlib import Path

import storage
from config import AppConfig

logger = logging.getLogger(__name__)

EXPORT_COLUMNS = [
    "id", "title", "url", "category", "source", "published_at", "collected_at",
    "status", "summary", "summary_model", "summarized_at", "sentiment", "sentiment_score", "content",
]


def _prepare(rows: list[dict]) -> list[dict]:
    prepared = []
    for row in rows:
        sentiment = {}
        if row.get("sentiment_result_json"):
            try:
                sentiment = json.loads(row["sentiment_result_json"])
            except ValueError:
                sentiment = {}
        prepared.append(
            {
                **{k: row.get(k) for k in EXPORT_COLUMNS if k in row},
                "status": "summarized" if row.get("summary") else "clean",
                "sentiment": sentiment.get("sentiment"),
                "sentiment_score": sentiment.get("score"),
            }
        )
    return prepared


def write_csv(rows: list[dict], path: Path) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=EXPORT_COLUMNS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_jsonl(rows: list[dict], path: Path) -> None:
    with path.open("w", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps({k: row.get(k) for k in EXPORT_COLUMNS}, ensure_ascii=False) + "\n")


def write_xlsx(rows: list[dict], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "news"
    ws.append(EXPORT_COLUMNS)
    for row in rows:
        ws.append([_cell(row.get(k)) for k in EXPORT_COLUMNS])
    widths = {"title": 40, "url": 40, "summary": 60, "content": 80}
    for index, name in enumerate(EXPORT_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = widths.get(name, 14)
    ws.freeze_panes = "A2"
    wb.save(path)


def _cell(value):
    if isinstance(value, str) and len(value) > 32000:  # Excel 셀 길이 제한
        return value[:32000]
    return value


WRITERS = {"csv": write_csv, "jsonl": write_jsonl, "xlsx": write_xlsx}


def run_export(
    config: AppConfig,
    *,
    fmt: str = "csv",
    status: str = "all",
    category: str | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
    out: str | None = None,
) -> str:
    """조건에 맞는 뉴스를 파일로 내보내고 경로를 반환한다."""
    if fmt not in WRITERS:
        raise ValueError(f"지원하지 않는 포맷: {fmt}")
    rows = _prepare(
        storage.export_rows(
            category=category, date_from=date_from, date_to=date_to,
            status=None if status == "all" else status,
        )
    )
    output_dir = config.resolve(config.report.get("output_dir", "output")) / "exports"
    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = Path(out) if out else output_dir / f"news_{status}_{datetime.now():%Y%m%d}.{fmt}"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    WRITERS[fmt](rows, out_path)
    logger.info("내보내기 완료: %s건 → %s", len(rows), out_path)
    return str(out_path)
